"""Regression coverage for the evidence-based historical review follow-up."""
import contextlib
import io
import runpy
import sys
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
from zipfile import ZipFile

import pytest
from docx import Document
from docxtpl import RichText

from sheetflow import common, personal_income_tax as pit, render
from sheetflow.named_ranges import build_named_range_map
from sheetflow.smart_path_manager import SmartPathManager


def test_real_word_batch_preserves_xml_characters_and_richtext(tmp_path):
    template = tmp_path / 'template.docx'
    document = Document()
    document.add_paragraph('{{ CN }}')
    document.add_paragraph('{{r rich }}')
    document.sections[0].header.paragraphs[0].text = '{{ CN }}'
    document.add_table(rows=1, cols=1).cell(0, 0).text = '{{ CN }}'
    document.save(template)
    records = [
        {'CN': 'A&B <公司>', 'rich': RichText('甲&乙 <丙>'), 'id': 1},
        {'CN': '第二家公司', 'rich': RichText('丁&戊'), 'id': 2},
    ]
    render.render_docx(records, template, tmp_path, filename=lambda r: str(r['id']) + '.docx')
    for record, rich in zip(records, ['甲&乙 <丙>', '丁&戊']):
        result = Document(tmp_path / (str(record['id']) + '.docx'))
        assert result.paragraphs[0].text == record['CN']
        assert result.paragraphs[1].text == rich
        assert result.tables[0].cell(0, 0).text == record['CN']
        assert result.sections[0].header.paragraphs[0].text == record['CN']


def test_actual_template_read_permission_is_formatted(monkeypatch, tmp_path):
    template = tmp_path / 'locked.docx'
    Document().save(template)
    real_open = io.open

    def deny_template_read(path, *args, **kwargs):
        if path == template or path == str(template):
            raise PermissionError('read denied')
        return real_open(path, *args, **kwargs)

    monkeypatch.setattr(io, 'open', deny_template_read)
    with pytest.raises(RuntimeError, match='ERR_PERMISSION_DENIED'):
        render.render_docx({'CN': 'Acme'}, template, tmp_path, filename=lambda r: 'result.docx')
    assert not (tmp_path / 'result.docx').exists()


def test_real_docxtpl_output_permission_is_not_template_error(monkeypatch, tmp_path):
    template = tmp_path / 'template.docx'
    Document().save(template)
    output = tmp_path / 'output.docx'
    real_open = io.open

    def deny_output(path, *args, **kwargs):
        if path == output or path == str(output):
            raise PermissionError('output denied')
        return real_open(path, *args, **kwargs)

    monkeypatch.setattr(io, 'open', deny_output)
    with pytest.raises(PermissionError, match='output denied'):
        render.render_docx({'CN': 'Acme'}, template, tmp_path, filename=lambda r: output.name)


@pytest.mark.parametrize('cn', [None, '', '   ', 123])
def test_tax_rejects_missing_company_before_directory_creation(cn, tmp_path):
    tax = runpy.run_path(Path(__file__).parents[1] / 'py_script' / '税务.py')
    create_output = Mock(return_value=tmp_path)
    with pytest.raises(ValueError, match='ERR_OUTPUT_PATH'):
        tax['_render_periodic_tax_docx'](
            {'CN': cn}, [(date(2026, 1, 1), date(2026, 1, 31))],
            tmp_path / 'template.docx', create_output,
        )
    create_output.assert_not_called()


def _tax_app(monkeypatch, tmp_path, *, save_error=None, close_error=None):
    book = Mock()
    book.sheets = [Mock()]
    book.save.side_effect = save_error
    book.close.side_effect = close_error
    app = SimpleNamespace(books=SimpleNamespace(open=Mock(return_value=book)))
    launch = Mock(return_value=contextlib.nullcontext(app))
    monkeypatch.setattr(pit.xw, 'App', launch)
    monkeypatch.setattr(pit, 'search_template_file_cached', lambda *a, **k: tmp_path / 'template.xls')
    return book, launch


@pytest.mark.parametrize('close_fails', [False, True])
def test_tax_close_does_not_mask_save_failure(monkeypatch, tmp_path, close_fails):
    save_error = RuntimeError('save failed')
    book, _ = _tax_app(monkeypatch, tmp_path, save_error=save_error,
                       close_error=RuntimeError('close failed') if close_fails else None)
    with pytest.raises(RuntimeError, match='save failed') as caught:
        pit.generate_personal_income_tax({}, [(date(2026, 1, 1), date(2026, 1, 31))], tmp_path)
    assert caught.value is save_error
    book.close.assert_called_once()


def test_tax_close_only_failure_still_propagates(monkeypatch, tmp_path):
    book, _ = _tax_app(monkeypatch, tmp_path, close_error=RuntimeError('close failed'))
    with pytest.raises(RuntimeError, match='close failed'):
        pit.generate_personal_income_tax({}, [(date(2026, 1, 1), date(2026, 1, 31))], tmp_path)
    book.save.assert_called_once()


def test_tax_duplicate_period_outputs_rejected_before_excel(monkeypatch, tmp_path):
    launch = Mock(side_effect=AssertionError('Excel must not start'))
    monkeypatch.setattr(pit.xw, 'App', launch)
    periods = iter([(date(2026, 1, 1), date(2026, 1, 15)),
                    (date(2026, 1, 16), date(2026, 1, 31))])
    with pytest.raises(ValueError, match='ERR_OUTPUT_PATH'):
        pit.generate_personal_income_tax({}, periods, tmp_path)
    launch.assert_not_called()
    assert list(tmp_path.iterdir()) == []


def test_tax_distinct_month_generator_still_generates_and_archives(monkeypatch, tmp_path):
    book, launch = _tax_app(monkeypatch, tmp_path)
    book.save.side_effect = lambda path: Path(path).write_text('generated', encoding='utf-8')
    periods = iter([(date(2018, 12, 1), date(2018, 12, 31)),
                    (date(2019, 1, 1), date(2019, 1, 31))])
    assert pit.generate_personal_income_tax({}, periods, tmp_path) == tmp_path
    assert launch.call_count == 1
    assert book.close.call_count == 2
    for stem, year in [('2018_12', '2018'), ('2019_01', '2019')]:
        with ZipFile(tmp_path / year / (stem + '.zip')) as archive:
            assert archive.read(stem + '.xls') == b'generated'
        assert not (tmp_path / (stem + '.xls')).exists()


@pytest.mark.parametrize('name', [123, 0, False, date(2026, 1, 1)])
def test_named_map_rejects_non_string_scalar_names(name):
    with pytest.raises(ValueError, match='ERR_NAMED_RANGE_MAP'):
        build_named_range_map({'Named': name, 'Range': 'A1'})


def test_named_map_rejects_non_string_duplicate_names():
    with pytest.raises(ValueError, match='ERR_NAMED_RANGE_MAP'):
        build_named_range_map({'Named': [123, 123], 'Range': ['A1', 'B1']})


@pytest.mark.parametrize('raw', ['', '.', 'relative', 'C:relative'])
def test_desktop_rejects_non_absolute_registry_values(monkeypatch, tmp_path, raw):
    monkeypatch.chdir(tmp_path)
    (tmp_path / 'relative').mkdir()
    fake = SimpleNamespace(
        HKEY_CURRENT_USER=object(),
        OpenKey=lambda *a: contextlib.nullcontext(),
        QueryValueEx=lambda *a: (raw, 1),
    )
    monkeypatch.setitem(sys.modules, 'winreg', fake)
    monkeypatch.setattr(sys, 'platform', 'win32')
    assert SmartPathManager._desktop_path() is None


@pytest.mark.parametrize('fullname', [None, '', 'Book1', 'relative/book.xlsx', 'C:book.xlsx'])
def test_workbook_path_invalid_before_creating_output(monkeypatch, fullname):
    manager = Mock(side_effect=AssertionError('Must validate before creation'))
    monkeypatch.setattr(common, 'SmartPathManager', manager)
    with pytest.raises(ValueError, match='ERR_INVALID_PATH'):
        common.create_output_dir_for_workbook(SimpleNamespace(fullname=fullname))
    manager.assert_not_called()


def test_workbook_directory_permission_has_output_error_code(monkeypatch, tmp_path):
    failure = PermissionError('directory denied')
    manager = Mock()
    manager.mkdir_if_needed.side_effect = failure
    monkeypatch.setattr(common, 'SmartPathManager', Mock(return_value=manager))
    with pytest.raises(RuntimeError, match='ERR_OUTPUT_DIR') as caught:
        common.create_output_dir_for_workbook(SimpleNamespace(fullname=str(tmp_path / 'book.xlsx')))
    assert caught.value.__cause__ is failure
